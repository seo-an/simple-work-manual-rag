"""Markdown sections + Excel file registry -> transactional SQLite index."""
from __future__ import annotations
from pathlib import Path
from contextlib import closing
import os
import sqlite3
import tempfile
from openpyxl import load_workbook

REQUIRED_COLUMNS = ("file_id", "title", "summary", "location", "extension", "keywords")


def parse_markdown(path: Path, root: Path) -> list[tuple[str, str, str, str]]:
    text = path.read_text(encoding="utf-8")
    doc_id = path.relative_to(root).as_posix()
    title = path.stem
    section = title
    buf: list[str] = []
    chunks = []

    def flush() -> None:
        body = "\n".join(buf).strip()
        if body:
            chunks.append((doc_id, title, section, body))
        buf.clear()

    for line in text.splitlines():
        if line.startswith("# "):
            flush()
            title = line[2:].strip()
            section = title
        elif line.startswith("## "):
            flush()
            section = line[3:].strip()
        else:
            buf.append(line)
    flush()
    return chunks


def read_registry(path: Path) -> list[tuple[str, str, str, str, str, str]]:
    if not path.is_file():
        raise FileNotFoundError(f"missing file registry: {path}")
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        columns = [str(c or "").strip() for c in next(rows)]
        if tuple(columns) != REQUIRED_COLUMNS:
            raise ValueError("registry columns must be " + ", ".join(REQUIRED_COLUMNS))
        result = []
        seen = set()
        for row in rows:
            if not any(value is not None and str(value).strip() for value in row):
                continue
            values = tuple(str(v or "").strip() for v in row[:len(REQUIRED_COLUMNS)])
            if not all(values[:5]) or values[0] in seen:
                raise ValueError("registry row has a missing field or duplicate file_id")
            seen.add(values[0])
            result.append(values)
        return result
    finally:
        book.close()


def build_index(data_dir: Path, db_path: Path) -> dict[str, int]:
    paths = sorted((data_dir / "manuals").rglob("*.md"))
    if not paths:
        raise ValueError("no Markdown manuals supplied")
    chunks = [item for path in paths for item in parse_markdown(path, data_dir)]
    rows = read_registry(data_dir / "file_registry.xlsx")
    db_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix="rag-index-", suffix=".db", dir=db_path.parent)
    os.close(fd)
    temp = Path(temp_name)
    try:
        with closing(sqlite3.connect(temp)) as conn:
            conn.executescript("""
                CREATE TABLE chunks (
                  id INTEGER PRIMARY KEY,
                  doc_id TEXT NOT NULL,
                  title TEXT NOT NULL,
                  section TEXT NOT NULL,
                  body TEXT NOT NULL
                );
                CREATE TABLE files (
                  file_id TEXT PRIMARY KEY,
                  title TEXT NOT NULL,
                  summary TEXT NOT NULL,
                  location TEXT NOT NULL,
                  extension TEXT NOT NULL,
                  keywords TEXT NOT NULL
                );
            """)
            conn.executemany("INSERT INTO chunks(doc_id,title,section,body) VALUES (?,?,?,?)", chunks)
            conn.executemany("INSERT INTO files VALUES (?,?,?,?,?,?)", rows)
            try:
                conn.execute("CREATE VIRTUAL TABLE chunks_fts USING fts5(body, title, tokenize='trigram')")
                conn.execute("INSERT INTO chunks_fts(rowid, body, title) SELECT id,body,title FROM chunks")
            except sqlite3.OperationalError:
                pass
            conn.commit()
        os.replace(temp, db_path)
    finally:
        temp.unlink(missing_ok=True)
    return {"manuals": len(paths), "chunks": len(chunks), "files": len(rows)}
